"""Palette, number formats and cell writers for the VRT V1 workbook."""
from __future__ import annotations

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

FONT = "Arial"

NAVY = "1F3864"
TEAL = "2E7D77"
BAND = "EAEFF7"
YELLOW = "FFF2CC"
GRAYFILL = "F2F2F2"
MISSING = "FCE4E4"
PROPOSED = "EFEFEF"
SECTION = "D9E1F2"

BLUE_TXT = "0000FF"
GREEN_TXT = "008000"
BLACK_TXT = "000000"
GRAY_TXT = "808080"
PINK_TXT = "C00000"

FMT_M = '#,##0.0;(#,##0.0);-'
FMT_PCT = '0.0%;(0.0%);-'
FMT_EPS = '0.00;(0.00);-'
FMT_X = '0.0x'
FMT_PRICE = '#,##0.00'
FMT_SHARES = '#,##0.0'
FMT_COUNT = '#,##0'
FMT_PP = '0.0 "pp";(0.0 "pp");-'

SEMANTIC_FORMATS = {
    "money": FMT_M,
    "percent": FMT_PCT,
    "percentage_points": FMT_PP,
    "eps": FMT_EPS,
    "multiple": FMT_X,
    "price": FMT_PRICE,
    "shares": FMT_SHARES,
    "count": FMT_COUNT,
    "text": "@",
}

KIND_FONT = {"hardcode": BLUE_TXT, "link": GREEN_TXT, "formula": BLACK_TXT,
             "override": BLUE_TXT, "gray": GRAY_TXT, "missing": PINK_TXT,
             "proposed": BLUE_TXT, "benchmark": BLUE_TXT, "text": BLACK_TXT}
KIND_FILL = {"override": YELLOW, "gray": GRAYFILL, "missing": MISSING, "proposed": PROPOSED}


def title(ws, text, last_col=26):
    cell = ws.cell(row=1, column=2, value=text)
    cell.font = Font(FONT, size=16, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    for col in range(3, last_col + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 24


def subtitle(ws, text):
    cell = ws.cell(row=3, column=2, value=text)
    cell.font = Font(FONT, size=10, color="595959")


def backlink(ws, text="Back to Outlook"):
    cell = ws.cell(row=4, column=2, value=text)
    cell.font = Font(FONT, size=10, color="0563C1", underline="single")


def header(ws, row, col, text, forecast=False, fill=None):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(FONT, size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=fill or (TEAL if forecast else NAVY))
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    return cell


def label(ws, row, text, bold=False, italic=False, color=BLACK_TXT, col=2):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(FONT, size=10, bold=bold, italic=italic, color=color)
    return cell


def note(ws, row, text, col=2):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(FONT, size=9, italic=True, color="595959")
    return cell


def band(ws, row, cols):
    ws.cell(row=row, column=2).fill = PatternFill("solid", fgColor=BAND)
    for col in cols:
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=BAND)


def put(ws, row, col, value, kind="formula", fmt=FMT_M, bold=False, wrap=False):
    """kind: hardcode | benchmark | link | formula | override | gray | missing | proposed | text"""
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(FONT, size=10, bold=bold, color=KIND_FONT[kind],
                     italic=kind == "proposed")
    if fmt:
        cell.number_format = fmt
    fill = KIND_FILL.get(kind)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def comment(cell, text, width=420, height=120):
    if not text:
        return cell
    cell.comment = Comment(text, "model build", width=width, height=height)
    return cell


def widths(ws, first_width=46, data=11, last_col=26, meta=None):
    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["B"].width = first_width
    for col in range(3, last_col + 1):
        ws.column_dimensions[L(col)].width = data
    for col, width in (meta or {}).items():
        ws.column_dimensions[L(col)].width = width


def base(ws, text, subtitle_text, last_col=26, first_width=46, backlink_text="Back to Outlook"):
    title(ws, text, last_col)
    subtitle(ws, subtitle_text)
    if backlink_text:
        backlink(ws, backlink_text)
    widths(ws, first_width, last_col=last_col)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C7"
