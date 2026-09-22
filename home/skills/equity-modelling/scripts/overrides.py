"""Extract and restore a model workbook's user-owned values.

Two things survive a rebuild, both by identity rather than by cell coordinate:

* **numeric assumption overrides**, identified by ``(driver, fiscal period)``; and
* **the designated Outlook text fields**, identified by their stable field id (located
  in the workbook by an exact, configured label).

Workbook contract (see the skill's ``references/model-rules.md``)
-----------------------------------------------------------------
The override sheet holds one four-row block per driver, rows adjacent and in order::

    <driver> | proposed (inactive)   — research, never referenced by a formula
    <driver> | approved default      — the user-approved house case
    <driver> | YOUR override         — blank until the user fills it in
    <driver> | active input          — formula selecting override, else approved

Only the override row carries the numeric identity; the driver name is its label with
the ``| YOUR override`` suffix removed.  Periods are labelled in a header row (default
row 6), one label per column, e.g. ``Q3 2026E``.  A blank override cell means "use the
default"; zero is a valid override.

The active-input formula is the driver's *definition*.  It is recorded with cell
references normalised to ``{d}`` / ``{o}`` / ``{a}`` (default, override, active row)
and the driver's own column, so a block that merely moves rows or columns keeps the
same definition signature while a changed calculation is detected and withheld.

The Outlook sheet carries the preserved text fields as ``label | text`` pairs: the
exact label in the label column (default ``A``) and the text in the text column
(default ``B``).  Each label must occur exactly once; a missing or duplicated label is
a hard error, because restoring by position would silently write to the wrong field.

Payload
-------
``schema_version`` 2, ``assumption_overrides`` (driver id → period → value implicitly,
as a list of records) and ``outlook_text`` (field id → text).  Version 1 payloads are
rejected: they carry no Outlook text and no stable driver ids.

CLI
---
    python3 overrides.py check   WB.xlsx
    python3 overrides.py extract WB.xlsx --out payload.json
    python3 overrides.py restore SRC.xlsx DST.xlsx --overrides payload.json [--report report.json]
"""
from __future__ import annotations

import argparse
from dataclasses import dataclass, field, asdict
import hashlib
import json
from pathlib import Path
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter as L

SCHEMA_VERSION = 2
OVERRIDE_SUFFIX = "YOUR override"
DEFAULT_SUFFIX = "default"
ACTIVE_SUFFIX = "active input"

PERIOD_RE = re.compile(r"^(?:Q([1-4])\s*(\d{4})|(\d{4})\s*Q([1-4]))\s*([AE])$", re.IGNORECASE)
REF_RE = re.compile(r"([A-Za-z]{1,3})(\d+)")

DEFAULT_OUTLOOK_FIELDS = {
    "thesis": "Thesis",
    "key_debate": "Key debate",
    "catalysts": "Catalysts",
    "risks": "Risks",
    "falsifying_evidence": "Falsifying evidence",
}


class OverrideError(Exception):
    """Incompatible workbook metadata, an invalid value, or an unsupported payload."""


def canonical_period(label: str) -> str:
    """``"Q3 2026E"`` -> ``"2026Q3E"``.  Idempotent; raises on any other shape."""
    match = PERIOD_RE.match(str(label).strip())
    if not match:
        raise OverrideError(f"Unrecognised fiscal period label: {label!r}")
    quarter, year, year2, quarter2, flag = match.groups()
    return f"{year or year2}Q{quarter or quarter2}{flag.upper()}"


def driver_slug(name: str) -> str:
    """A stable, readable id for a driver label (used when the caller supplies none)."""
    slug = re.sub(r"[^a-z0-9]+", "_", str(name).strip().casefold()).strip("_")
    return slug or "driver"


@dataclass(frozen=True)
class Driver:
    name: str
    driver_id: str
    default_row: int
    override_row: int
    active_row: int
    number_format: str


@dataclass(frozen=True)
class Override:
    driver: str
    driver_id: str
    period: str
    value: float
    cell: str
    number_format: str
    definition: str


@dataclass
class OverrideSet:
    sheet: str = "Assumptions"
    header_row: int = 6
    label_col: int = 2
    periods: dict[int, str] = field(default_factory=dict)
    drivers: dict[str, Driver] = field(default_factory=dict)
    overrides: dict[str, Override] = field(default_factory=dict)
    outlook_text: dict[str, str] = field(default_factory=dict)
    outlook_fields: dict[str, str] = field(default_factory=dict)
    outlook_sheet: str = "Outlook"
    outlook_label_col: int = 1
    outlook_text_col: int = 2
    source: str = ""

    def key(self, driver_id: str, period: str) -> str:
        return f"{driver_id}\u0000{canonical_period(period)}"

    def get(self, driver: str, period: str) -> Override | None:
        """Look up by driver label or driver id and period label."""
        driver_id = self.drivers[driver].driver_id if driver in self.drivers else driver
        return self.overrides.get(self.key(driver_id, period))

    def save(self, path: str | Path) -> Path:
        payload = {
            "schema_version": SCHEMA_VERSION,
            "sheet": self.sheet,
            "header_row": self.header_row,
            "label_col": self.label_col,
            "outlook_sheet": self.outlook_sheet,
            "outlook_label_col": self.outlook_label_col,
            "outlook_text_col": self.outlook_text_col,
            "source": self.source,
            "periods": {str(col): period for col, period in self.periods.items()},
            "drivers": {name: asdict(driver) for name, driver in self.drivers.items()},
            "assumption_overrides": [asdict(item) for item in self.overrides.values()],
            "outlook_text": dict(self.outlook_text),
            "outlook_fields": dict(self.outlook_fields),
        }
        path = Path(path)
        path.write_text(json.dumps(payload, indent=2) + "\n")
        return path

    @classmethod
    def load(cls, path: str | Path) -> "OverrideSet":
        payload = json.loads(Path(path).read_text())
        version = payload.get("schema_version")
        if version != SCHEMA_VERSION:
            raise OverrideError(f"{Path(path).name}: unsupported payload schema_version "
                                f"{version!r} (this helper writes and reads {SCHEMA_VERSION})")
        result = cls(
            sheet=payload["sheet"],
            header_row=payload["header_row"],
            label_col=payload["label_col"],
            outlook_sheet=payload.get("outlook_sheet", "Outlook"),
            outlook_label_col=payload.get("outlook_label_col", 1),
            outlook_text_col=payload.get("outlook_text_col", 2),
            source=payload.get("source", ""),
        )
        result.periods = {int(col): period for col, period in payload["periods"].items()}
        result.drivers = {name: Driver(**item) for name, item in payload["drivers"].items()}
        result.overrides = {}
        for item in payload["assumption_overrides"]:
            for name, driver in result.drivers.items():
                if driver.driver_id == item["driver_id"] and name != item["driver"]:
                    raise OverrideError(f"payload maps driver_id {item['driver_id']!r} to both "
                                        f"{name!r} and {item['driver']!r}")
            if not any(driver.driver_id == item["driver_id"] for driver in result.drivers.values()):
                raise OverrideError(f"payload carries an unknown driver id {item['driver_id']!r}")
            key = result.key(item["driver_id"], item["period"])
            if key in result.overrides:
                raise OverrideError(f"payload carries a duplicate override "
                                    f"{item['driver_id']} / {item['period']}")
            result.overrides[key] = Override(**item)
        result.outlook_fields = payload.get("outlook_fields", {})
        result.outlook_text = payload.get("outlook_text", {})
        for field_id in result.outlook_text:
            if result.outlook_fields and field_id not in result.outlook_fields:
                raise OverrideError(f"payload carries unknown Outlook field id {field_id!r}")
        return result


@dataclass
class RestoreReport:
    restored: list[Override] = field(default_factory=list)
    retired_period: list[Override] = field(default_factory=list)
    removed_driver: list[Override] = field(default_factory=list)
    definition_changed: list[Override] = field(default_factory=list)
    definition_change_allowed: list[Override] = field(default_factory=list)
    outlook_restored: dict[str, str] = field(default_factory=dict)
    outlook_missing: list[str] = field(default_factory=list)
    original_unchanged: bool = False
    destination: str = ""

    @property
    def needs_review(self) -> list[Override]:
        return self.definition_changed

    def lines(self) -> list[str]:
        out = [
            f"{len(self.restored)} override(s) restored, {len(self.retired_period)} retired period(s), "
            f"{len(self.removed_driver)} removed driver(s), "
            f"{len(self.definition_changed)} changed definition(s) "
            f"({len(self.definition_change_allowed)} restored as an allowed change), "
            f"{len(self.outlook_restored)} Outlook field(s) restored"
            + (f", {len(self.outlook_missing)} Outlook label(s) missing"
               if self.outlook_missing else "")
        ]
        for item in self.retired_period:
            out.append(f"  retired (period no longer in the forecast grid): {item.driver} "
                       f"{item.period} = {item.value}")
        for item in self.removed_driver:
            out.append(f"  removed driver (not restored): {item.driver} {item.period} = {item.value}")
        for item in self.definition_changed:
            marker = "CHANGED DEFINITION" + (" (allowed)" if item in self.definition_change_allowed
                                             else ", review required")
            out.append(f"  {marker}: {item.driver} {item.period} = {item.value}")
        for field_id in self.outlook_missing:
            out.append(f"  Outlook label missing in the rebuild: {field_id!r} (nothing written)")
        out.append(f"  original workbook unchanged: {self.original_unchanged}")
        return out

    def to_json(self, path: str | Path) -> Path:
        payload = {
            "destination": self.destination,
            "original_unchanged": self.original_unchanged,
            "restored": [asdict(item) for item in self.restored],
            "retired_period": [asdict(item) for item in self.retired_period],
            "removed_driver": [asdict(item) for item in self.removed_driver],
            "definition_changed": [asdict(item) for item in self.definition_changed],
            "definition_change_allowed": [asdict(item) for item in self.definition_change_allowed],
            "outlook_restored": self.outlook_restored,
            "outlook_missing": self.outlook_missing,
        }
        path = Path(path)
        path.write_text(json.dumps(payload, indent=2) + "\n")
        return path


def sha256(path: str | Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _suffix_match(text: str, suffix: str) -> str | None:
    """Return the driver name if ``text`` ends with the suffix (case-insensitive)."""
    if not isinstance(text, str):
        return None
    stripped = text.strip()
    if stripped.casefold().endswith(suffix.strip().casefold()):
        return stripped[: len(stripped) - len(suffix.strip())].strip().rstrip("|").strip()
    return None


def _signature(formula, col_letter: str, rows: dict[int, str]) -> str:
    """Normalise a driver's active-input formula so row/column moves do not matter."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return str(formula)

    def replace(match: re.Match) -> str:
        col, row = match.group(1).upper(), int(match.group(2))
        if col == col_letter.upper() and row in rows:
            return "{%s}" % rows[row]
        return match.group(0)

    return REF_RE.sub(replace, formula)


def _numeric(cell, context: str) -> float:
    if isinstance(cell.value, bool):
        raise OverrideError(f"{context}: boolean is not a numeric override")
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        raise OverrideError(f"{context}: formulas are not valid overrides")
    if not isinstance(cell.value, (int, float)):
        raise OverrideError(f"{context}: {cell.value!r} is not numeric")
    return float(cell.value)


def _read_text(ws, row: int, col: int) -> str:
    """The text of a field, reading the merged top-left cell when the row is merged."""
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= col <= merged.max_col:
            value = ws.cell(merged.min_row, merged.min_col).value
            return "" if value is None else str(value)
    value = ws.cell(row, col).value
    return "" if value is None else str(value)


def extract_outlook(workbook, sheet: str, label_col: int, text_col: int,
                    fields: dict[str, str]) -> dict[str, str]:
    """Read the designated Outlook text fields by exact label."""
    if sheet not in workbook.sheetnames:
        raise OverrideError(f"no sheet named {sheet!r} (found {workbook.sheetnames})")
    ws = workbook[sheet]
    found: dict[str, str] = {}
    for field_id, label in fields.items():
        rows = [row for row in range(1, ws.max_row + 1)
                if ws.cell(row, label_col).value is not None
                and str(ws.cell(row, label_col).value).strip() == label]
        if len(rows) != 1:
            raise OverrideError(f"{sheet}!{L(label_col)}: label {label!r} occurs {len(rows)} times "
                                f"(expected exactly one)")
        found[field_id] = _read_text(ws, rows[0], text_col)
    return found


def extract(path: str | Path, sheet: str = "Assumptions", header_row: int = 6,
            label_col: int = 2, driver_ids: dict[str, str] | None = None,
            outlook_sheet: str = "Outlook", outlook_label_col: int = 1, outlook_text_col: int = 2,
            outlook_fields: dict[str, str] | None = None) -> OverrideSet:
    """Read every designated override and Outlook text field without modifying the workbook."""
    path = Path(path)
    workbook = load_workbook(path, data_only=False)
    if sheet not in workbook.sheetnames:
        raise OverrideError(f"{path.name}: no sheet named {sheet!r} (found {workbook.sheetnames})")
    ws = workbook[sheet]

    periods: dict[int, str] = {}
    for col in range(label_col + 1, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        if value is None or str(value).strip() == "":
            continue
        periods[col] = canonical_period(value)
    if not periods:
        raise OverrideError(f"{path.name}: no period labels in row {header_row} of {sheet}")
    if len(set(periods.values())) != len(periods):
        raise OverrideError(f"{path.name}: duplicate period labels in row {header_row}: "
                            f"{list(periods.values())}")

    drivers: dict[str, Driver] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        name = _suffix_match(ws.cell(row, label_col).value, OVERRIDE_SUFFIX)
        if not name:
            continue
        if name in drivers:
            raise OverrideError(f"{path.name}: duplicate override row for driver {name!r}")
        default_row, active_row = row - 1, row + 1
        if not _suffix_match(ws.cell(default_row, label_col).value, DEFAULT_SUFFIX):
            raise OverrideError(f"{path.name}: row {default_row} above {name!r} is not a default row")
        if not _suffix_match(ws.cell(active_row, label_col).value, ACTIVE_SUFFIX):
            raise OverrideError(f"{path.name}: row {active_row} below {name!r} is not an active-input row")
        drivers[name] = Driver(name=name, driver_id=(driver_ids or {}).get(name) or driver_slug(name),
                               default_row=default_row, override_row=row, active_row=active_row,
                               number_format=ws.cell(row, min(periods)).number_format)
    if not drivers:
        raise OverrideError(f"{path.name}: no '{OVERRIDE_SUFFIX}' driver rows found on {sheet}")
    if driver_ids:
        unknown = sorted(set(driver_ids) - set(drivers))
        if unknown:
            raise OverrideError(f"{path.name}: driver id map names drivers that are not in the "
                                f"workbook: {unknown}")
    if len({driver.driver_id for driver in drivers.values()}) != len(drivers):
        raise OverrideError(f"{path.name}: two drivers resolve to the same driver id")

    result = OverrideSet(sheet=sheet, header_row=header_row, label_col=label_col,
                         periods=periods, drivers=drivers, source=str(path),
                         outlook_sheet=outlook_sheet, outlook_label_col=outlook_label_col,
                         outlook_text_col=outlook_text_col,
                         outlook_fields=dict(outlook_fields or DEFAULT_OUTLOOK_FIELDS))
    for name, driver in drivers.items():
        rows = {driver.default_row: "d", driver.override_row: "o", driver.active_row: "a"}
        for col, period in periods.items():
            letter = L(col)
            cell = ws.cell(driver.override_row, col)
            definition = _signature(ws.cell(driver.active_row, col).value, letter, rows)
            if cell.value is None:
                continue
            key = result.key(driver.driver_id, period)
            if key in result.overrides:
                raise OverrideError(f"{path.name}: duplicate override identity {name!r} / {period}")
            context = f"{path.name}: {sheet}!{cell.coordinate} ({name} / {period})"
            result.overrides[key] = Override(
                driver=name, driver_id=driver.driver_id, period=period,
                value=_numeric(cell, context), cell=cell.coordinate,
                number_format=cell.number_format or driver.number_format, definition=definition)
    if outlook_sheet in workbook.sheetnames:
        result.outlook_text = extract_outlook(workbook, outlook_sheet, outlook_label_col,
                                              outlook_text_col, result.outlook_fields)
    else:
        result.outlook_text = {}          # no Outlook sheet: nothing to preserve

    return result


def restore(template: str | Path, destination: str | Path, overrides: OverrideSet,
            allow_definition_change: bool = False) -> RestoreReport:
    """Write ``template`` to ``destination`` with the recorded overrides restored.

    ``template`` is the freshly rebuilt workbook (new grid, current defaults); the
    workbook the payload came from is never opened.  Placement comes from the
    template's own period header, driver rows and Outlook labels, so a rolled-forward
    grid, a moved block or a reordered Outlook section are handled by identity.
    """
    template, destination = Path(template).resolve(), Path(destination).resolve()
    if template == destination:
        raise OverrideError("refusing to write over the template workbook; pass a separate destination")
    before = sha256(template)

    workbook = load_workbook(template, data_only=False)
    if overrides.sheet not in workbook.sheetnames:
        raise OverrideError(f"{template.name}: no sheet named {overrides.sheet!r}")
    ws = workbook[overrides.sheet]

    periods: dict[int, str] = {}
    for col in range(overrides.label_col + 1, ws.max_column + 1):
        value = ws.cell(overrides.header_row, col).value
        if value is None or str(value).strip() == "":
            continue
        periods[col] = canonical_period(value)
    if not periods:
        raise OverrideError(f"{template.name}: no period labels in row {overrides.header_row}")
    if len(set(periods.values())) != len(periods):
        raise OverrideError(f"{template.name}: duplicate period labels in row {overrides.header_row}: "
                            f"{list(periods.values())}")
    by_period = {period: col for col, period in periods.items()}

    rows_by_driver: dict[str, int] = {}
    for row in range(overrides.header_row + 1, ws.max_row + 1):
        name = _suffix_match(ws.cell(row, overrides.label_col).value, OVERRIDE_SUFFIX)
        if not name:
            continue
        if name in rows_by_driver:
            raise OverrideError(f"{template.name}: duplicate override row for driver {name!r}")
        rows_by_driver[name] = row

    report = RestoreReport(destination=str(destination))
    for item in overrides.overrides.values():
        if item.period not in by_period:
            report.retired_period.append(item)
            continue
        row = rows_by_driver.get(item.driver)
        if row is None:
            report.removed_driver.append(item)
            continue
        col = by_period[item.period]
        driver = overrides.drivers[item.driver]
        # The block may have moved rows: rebuild the row map around its new location
        # so a moved block keeps its definition signature.
        shift = row - driver.override_row
        rows = {driver.default_row + shift: "d", row: "o", driver.active_row + shift: "a"}
        current = _signature(ws.cell(driver.active_row + shift, col).value, L(col), rows)
        if current != item.definition:
            report.definition_changed.append(item)
            if not allow_definition_change:
                continue
            report.definition_change_allowed.append(item)
        cell = ws.cell(row, col)
        cell.value = item.value
        cell.number_format = item.number_format
        report.restored.append(item)

    if overrides.outlook_text:
        if overrides.outlook_sheet not in workbook.sheetnames:
            report.outlook_missing.extend(sorted(overrides.outlook_text))
        else:
            outlook = workbook[overrides.outlook_sheet]
            for field_id, text in overrides.outlook_text.items():
                label = overrides.outlook_fields.get(field_id, field_id)
                rows = [row for row in range(1, outlook.max_row + 1)
                        if outlook.cell(row, overrides.outlook_label_col).value is not None
                        and str(outlook.cell(row, overrides.outlook_label_col).value).strip() == label]
                if len(rows) != 1:
                    raise OverrideError(f"{template.name}: Outlook label {label!r} occurs {len(rows)} "
                                        f"times (expected exactly one)")
                target = outlook.cell(rows[0], overrides.outlook_text_col)
                target.value = text if text != "" else None
                report.outlook_restored[field_id] = text

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    report.original_unchanged = sha256(template) == before
    if not report.original_unchanged:
        raise OverrideError(f"{template} changed during restore; the output may be unusable")
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("command", choices=["check", "extract", "restore"])
    parser.add_argument("workbook", nargs="?",
                        help="workbook to read; for restore, the freshly rebuilt template")
    parser.add_argument("destination", nargs="?", help="restore output path (template is never written)")
    parser.add_argument("--sheet", default="Assumptions")
    parser.add_argument("--header-row", type=int, default=6)
    parser.add_argument("--outlook-sheet", default="Outlook")
    parser.add_argument("--outlook-label-col", default="A")
    parser.add_argument("--outlook-text-col", default="B")
    parser.add_argument("--overrides", help="payload JSON written by `extract`")
    parser.add_argument("--out", help="write the extracted payload here")
    parser.add_argument("--report", help="write the restore report JSON here")
    parser.add_argument("--allow-definition-change", action="store_true",
                        help="restore overrides whose driver definition changed (they are still "
                             "reported as an allowed definition change)")
    args = parser.parse_args(argv)
    label_col = column_index_from_string(args.outlook_label_col) if args.outlook_label_col.isalpha() \
        else int(args.outlook_label_col)
    text_col = column_index_from_string(args.outlook_text_col) if args.outlook_text_col.isalpha() \
        else int(args.outlook_text_col)

    if args.command == "check":
        found = extract(args.workbook, args.sheet, args.header_row, outlook_sheet=args.outlook_sheet,
                        outlook_label_col=label_col, outlook_text_col=text_col)
        print(f"{len(found.drivers)} drivers, periods {sorted(set(found.periods.values()))}")
        print(f"{len(found.overrides)} override(s):")
        for item in found.overrides.values():
            print(f"  {item.driver_id} {item.period} = {item.value} @ {item.cell} [{item.number_format}]")
        print(f"{len(found.outlook_text)} Outlook field(s): "
              f"{sorted(found.outlook_text)}")
        for field_id, text in found.outlook_text.items():
            preview = text if len(text) <= 60 else text[:57] + "..."
            print(f"  {field_id}: {preview!r}")
        return 0

    if args.command == "extract":
        found = extract(args.workbook, args.sheet, args.header_row, outlook_sheet=args.outlook_sheet,
                        outlook_label_col=label_col, outlook_text_col=text_col)
        out = Path(args.out or "overrides.json")
        found.save(out)
        print(f"{len(found.overrides)} override(s) and {len(found.outlook_text)} Outlook field(s) "
              f"written to {out}")
        return 0

    if not args.overrides or not args.destination:
        parser.error("restore needs DESTINATION and --overrides")
    report = restore(args.workbook, args.destination, OverrideSet.load(args.overrides),
                     allow_definition_change=args.allow_definition_change)
    for line in report.lines():
        print(line)
    if args.report:
        report.to_json(args.report)
    return 0


if __name__ == "__main__":
    sys.exit(main())
