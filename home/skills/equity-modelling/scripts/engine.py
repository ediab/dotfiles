"""Small generic workbook spine for evidence-led equity models.

The engine owns period placement, approval activation and named references.  Company modules
own model-specific rows and formulas; they receive metric-keyed maps rather than cell addresses.
"""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
import json
from typing import Any, Iterable, Mapping, Protocol

from openpyxl import Workbook
from openpyxl.styles import PatternFill

try:  # Supports both `python scripts/engine.py` consumers and package imports.
    from .grid import KeyedReferences, PeriodGrid, labels
    from .style import FMT_M, SEMANTIC_FORMATS, base, header, label, put
except ImportError:  # pragma: no cover - exercised by company-project entry points.
    from grid import KeyedReferences, PeriodGrid, labels
    from style import FMT_M, SEMANTIC_FORMATS, base, header, label, put


class EngineError(ValueError):
    """A company project violates the generic workbook contract."""


class CompanyModule(Protocol):
    def workbook_rows(self, actuals: Mapping[str, Mapping[str, float | None]],
                      drivers: Mapping[str, Mapping[str, float | None]],
                      periods: PeriodGrid) -> Mapping[str, Mapping[str, Mapping[str, Any]]]: ...


@dataclass(frozen=True)
class BuildContext:
    workbook: Workbook
    periods: PeriodGrid
    references: KeyedReferences
    actuals: Mapping[str, Mapping[str, float | None]]
    active_drivers: Mapping[str, Mapping[str, float | None]]
    module_rows: Mapping[str, Mapping[str, Any]]


def _forecast_periods(spec: Mapping[str, Any]) -> tuple[str, ...]:
    raw = tuple(spec["periods"]["forecast_quarters"])
    result = tuple(period if period.endswith("E") else period + "E" for period in raw)
    if len(set(result)) != len(result):
        raise EngineError("duplicate forecast periods in model_spec.json")
    return result


def _source_period(period: str) -> str:
    return period[:-1] if period.endswith("E") else period


def _forecast_gate(spec: Mapping[str, Any]) -> bool | None:
    """Return the declared `forecast_gate.approved` state, or `None` when undeclared.

    A declared gate must be an explicit boolean so an ambiguous value can never silently
    activate a forecast.  An undeclared gate is treated as not approved: activation then
    fails closed rather than defaulting to active.
    """
    gate = spec.get("forecast_gate")
    if gate is None:
        return None
    if not isinstance(gate, Mapping):
        raise EngineError("model_spec.json forecast_gate must be an object")
    approved = gate.get("approved")
    if not isinstance(approved, bool):
        raise EngineError("model_spec.json forecast_gate.approved must be a boolean")
    return approved


def approval_digest(triples: Iterable[tuple[str, str, float | int]]) -> str:
    """Hash exact driver/period/value triples in a stable, reviewable representation."""
    normalized = sorted((driver, period, float(value)) for driver, period, value in triples)
    payload = json.dumps(normalized, separators=(",", ":"), ensure_ascii=True)
    return sha256(payload.encode()).hexdigest()


def _approval_hashes(manifest: Mapping[str, Any]) -> Mapping[str, str]:
    approvals = manifest.get("approvals", manifest)
    if not isinstance(approvals, Mapping):
        raise EngineError("approval manifest must be an object keyed by approval_id")
    hashes: dict[str, str] = {}
    for approval_id, entry in approvals.items():
        if isinstance(entry, str):
            digest = entry
        elif isinstance(entry, Mapping):
            digest = entry.get("sha256", entry.get("hash"))
        else:
            raise EngineError(f"approval manifest entry {approval_id!r} is invalid")
        if not isinstance(digest, str) or len(digest) != 64:
            raise EngineError(f"approval manifest entry {approval_id!r} lacks a SHA-256 hash")
        hashes[str(approval_id)] = digest.lower()
    return hashes


def active_driver_values(drivers: Iterable[Mapping[str, Any]], forecast_periods: Iterable[str],
                         approval_manifest: Mapping[str, Any] | None,
                         *, forecast_gate_approved: bool | None = None) -> dict[str, dict[str, float | None]]:
    """Return driver values only for hash-approved complete rows whose forecast gate is open.

    Proposed rows intentionally return blanks.  Approved rows activate only when
    `forecast_gate_approved` is exactly `True`: an undeclared gate (`None`) or a closed gate
    (`False`) fails closed, so a manifest alone can never switch on an unapproved forecast.
    A partial or stale approval is an error rather than a silent partial forecast.
    """
    periods = tuple(forecast_periods)
    values: dict[str, dict[str, float | None]] = {}
    approved: list[Mapping[str, Any]] = []
    for driver in drivers:
        driver_id = str(driver.get("driver_id", ""))
        if not driver_id or driver_id in values:
            raise EngineError(f"driver_id must be unique and non-empty: {driver_id!r}")
        row: dict[str, float | None] = {}
        for period in periods:
            raw = driver.get(_source_period(period))
            if raw in (None, ""):
                row[period] = None
            else:
                try:
                    row[period] = float(raw)
                except (TypeError, ValueError) as exc:
                    raise EngineError(f"{driver_id} {_source_period(period)} is not numeric") from exc
        values[driver_id] = row
        if driver.get("status") == "approved":
            approved.append(driver)
        elif driver.get("status") != "proposed":
            raise EngineError(f"{driver_id} has invalid status {driver.get('status')!r}")

    if not approved:
        return {driver_id: {period: None for period in periods} for driver_id in values}
    if len(approved) != len(values):
        raise EngineError("partial forecast approval activates nothing")
    if forecast_gate_approved is None:
        raise EngineError(
            "model_spec.json must declare forecast_gate.approved before approved forecast inputs can activate")
    if not forecast_gate_approved:
        raise EngineError("forecast_gate.approved is false; approved forecast inputs stay inactive")
    if approval_manifest is None:
        raise EngineError("approved drivers require an approval manifest")

    approval_ids = {str(driver.get("approval_id", "")) for driver in approved}
    if len(approval_ids) != 1 or "" in approval_ids:
        raise EngineError("all approved drivers must share one approval_id")
    expected = _approval_hashes(approval_manifest).get(approval_ids.pop())
    if expected is None:
        raise EngineError("approved driver batch is absent from the approval manifest")
    triples = [(driver_id, _source_period(period), value)
               for driver_id, row in values.items() for period, value in row.items()
               if value is not None]
    if len(triples) != len(values) * len(periods):
        raise EngineError("an approved batch may not contain blank forecast values")
    if approval_digest(triples) != expected:
        raise EngineError("driver values do not match their approved manifest hash")
    return values


def _actual_map(actuals: Iterable[Mapping[str, Any]], historical_periods: Iterable[str]) -> dict[str, dict[str, float | None]]:
    periods = set(historical_periods)
    result: dict[str, dict[str, float | None]] = {}
    seen: set[tuple[str, str]] = set()
    for fact in actuals:
        metric, period = str(fact.get("metric", "")), str(fact.get("period", ""))
        if not metric or period not in periods:
            continue
        key = metric, period
        if key in seen:
            raise EngineError(f"duplicate actual fact {metric} {period}")
        seen.add(key)
        raw = fact.get("value")
        value = None if raw in (None, "") else float(raw)
        result.setdefault(metric, {})[period] = value
    return result


NINE_SHEET_ORDER = ("Outlook", "Operating Model", "Bridge Model", "Financial Statements",
                    "Valuation", "Inputs", "Consensus", "SourceData", "Checks")


def _human_label(metric: str) -> str:
    """Readable fallback for stable internal metric IDs."""
    words = metric.replace("_", " ").split()
    acronyms = {"gaap": "GAAP", "eps": "EPS", "fx": "FX", "yoy": "YoY",
                "qoq": "QoQ", "ebitda": "EBITDA", "ebit": "EBIT", "ev": "EV"}
    return " ".join(acronyms.get(word.lower(), word[:1].upper() + word[1:]) for word in words)


def _write_actuals(ws, data: Mapping[str, Mapping[str, float | None]], grid: PeriodGrid,
                   references: KeyedReferences, price: Mapping[str, Any] | None) -> None:
    price_column = grid.first_column + len(grid.actuals) if price else None
    base(ws, "Source actuals", "Frozen canonical facts; unavailable is blank, never zero.",
         last_col=price_column or 2 + len(grid.actuals), first_width=34, backlink_text=None)
    header(ws, 6, 2, "Metric")
    for period in grid.actuals:
        header(ws, 6, grid.columns[period], labels([period])[0])
    if price_column:
        header(ws, 6, price_column, "Dated benchmark")
    for row, metric in enumerate(sorted(data), start=7):
        label(ws, row, _human_label(metric))
        for period in grid.actuals:
            value = data[metric].get(period)
            put(ws, row, grid.columns[period], value, "missing" if value is None else "hardcode", FMT_M)
            references.bind(f"actual.{metric}.{period}", ws.title, row, grid.columns[period])
    if price:
        try:
            value = float(price["value"])
        except (KeyError, TypeError, ValueError) as exc:
            raise EngineError("model_spec.json price.value must be numeric") from exc
        row = 7 + len(data)
        label(ws, row, f"price_dated ({price.get('date', 'undated')})")
        put(ws, row, price_column, value, "hardcode", FMT_M)
        references.bind("benchmark.price_dated", ws.title, row, price_column)


def _write_drivers(ws, drivers: Iterable[Mapping[str, Any]], active: Mapping[str, Mapping[str, float | None]],
                   grid: PeriodGrid, references: KeyedReferences) -> None:
    base(ws, "Forecast drivers", "Proposed inputs do not activate; approved values are manifest-hash bound.",
         last_col=3 + len(grid.forecasts), first_width=32, backlink_text=None)
    header(ws, 6, 2, "Driver")
    header(ws, 6, 3, "Status")
    for period in grid.forecasts:
        header(ws, 6, grid.columns[period] + 1, labels([period])[0], forecast=True)
    for row, driver in enumerate(drivers, start=7):
        driver_id = str(driver["driver_id"])
        label(ws, row, str(driver.get("driver_name", driver_id)))
        put(ws, row, 3, driver.get("status"), "text", "@")
        for period in grid.forecasts:
            value = active[driver_id][period]
            put(ws, row, grid.columns[period] + 1, value, "missing" if value is None else "hardcode", FMT_M)
            references.bind(f"driver.{driver_id}.{period}", ws.title, row, grid.columns[period] + 1)


def _module_rows(module: CompanyModule, actuals, drivers, grid: PeriodGrid) -> Mapping[str, Mapping[str, Mapping[str, Any]]]:
    rows = module.workbook_rows(actuals, drivers, grid)
    if not isinstance(rows, Mapping):
        raise EngineError("company module must return a sheet-keyed mapping")
    return rows


def _row_parts(metric: str, row: Any, sheet: str):
    """Accept legacy period maps or period maps with optional presentation metadata."""
    if not isinstance(row, Mapping):
        raise EngineError(f"{sheet} {metric} output must be a mapping")
    metadata = row.get("_meta", {})
    if not isinstance(metadata, Mapping):
        raise EngineError(f"{sheet} {metric} _meta must be a mapping")
    allowed = {"label", "format", "section", "order"}
    unknown_metadata = set(metadata) - allowed
    if unknown_metadata:
        raise EngineError(f"{sheet} {metric} has unsupported row metadata {sorted(unknown_metadata)}")
    unknown_format = metadata.get("format", "money") not in SEMANTIC_FORMATS
    if unknown_format:
        raise EngineError(f"{sheet} {metric} has unsupported semantic format {metadata.get('format')!r}")
    values = {key: value for key, value in row.items() if key != "_meta"}
    return values, metadata


def _write_module_rows(ws, rows: Mapping[str, Any], grid: PeriodGrid,
                       references: KeyedReferences, forecast_active: bool,
                       legacy_order: bool = False) -> None:
    """Render keyed company rows; optional `_meta` supplies label, format, section and order."""
    if not isinstance(rows, Mapping):
        raise EngineError(f"{ws.title} module output must be a metric-keyed mapping")
    base(ws, ws.title, "Company formulas use metric-keyed defined names only.",
         last_col=max(3, 2 + len(grid.periods)), first_width=34, backlink_text=None)
    header(ws, 6, 2, "Metric")
    for period in grid.periods:
        header(ws, 6, grid.columns[period], labels([period])[0], forecast=period.endswith("E"))
    prepared = []
    for insertion_order, (metric, raw_row) in enumerate(rows.items()):
        if not isinstance(metric, str) or not metric:
            raise EngineError(f"{ws.title} output rows must have non-empty metric ids")
        values, metadata = _row_parts(metric, raw_row, ws.title)
        unknown = set(values) - set(grid.periods)
        if unknown:
            raise EngineError(f"{ws.title} {metric} has unknown periods {sorted(unknown)}")
        order = metadata.get("order", insertion_order)
        if not isinstance(order, (int, float)):
            raise EngineError(f"{ws.title} {metric} row order must be numeric")
        prepared.append((order, insertion_order, metric, values, metadata))
    if legacy_order:
        prepared.sort(key=lambda item: item[2])
    else:
        prepared.sort(key=lambda item: (item[0], item[1]))
    row_number = 7
    previous_section = None
    for _order, _insertion_order, metric, values, metadata in prepared:
        section = metadata.get("section")
        if section and section != previous_section:
            label(ws, row_number, str(section), bold=True)
            for column in range(2, 3 + len(grid.periods)):
                ws.cell(row=row_number, column=column).fill = PatternFill("solid", fgColor="D9E1F2")
            row_number += 1
        previous_section = section
        row = row_number
        row_number += 1
        label(ws, row, str(metadata.get("label", _human_label(metric))))
        fmt = SEMANTIC_FORMATS[metadata.get("format", "money")]
        for period in grid.periods:
            proposed_value = values.get(period)
            if isinstance(proposed_value, str) and proposed_value.startswith("=") and "!" in proposed_value:
                raise EngineError(f"{ws.title} {metric} {period} uses a positional cross-sheet formula")
            value = proposed_value if (not period.endswith("E") or forecast_active) else None
            kind = "formula" if isinstance(value, str) and value.startswith("=") else \
                ("missing" if value is None else "hardcode")
            put(ws, row, grid.columns[period], value, kind, fmt)
            references.bind(f"model.{ws.title}.{metric}.{period}", ws.title, row, grid.columns[period])


def _write_inputs(ws, drivers: Iterable[Mapping[str, Any]], active,
                  grid: PeriodGrid, references: KeyedReferences) -> None:
    """Render editable forecast proposals separately from formula-bearing company sheets."""
    last_col = 3 + len(grid.forecasts)
    base(ws, "Inputs", "Proposed assumptions are visible for review but do not feed forecasts.",
         last_col=max(4, last_col), first_width=34, backlink_text=None)
    header(ws, 6, 2, "Assumption")
    header(ws, 6, 3, "Status")
    for index, period in enumerate(grid.forecasts, start=4):
        header(ws, 6, index, labels([period])[0], forecast=True)
    for row_number, driver in enumerate(drivers, start=7):
        driver_id = str(driver["driver_id"])
        label(ws, row_number, str(driver.get("driver_name", _human_label(driver_id))))
        status = "Approved" if driver.get("status") == "approved" else "Proposed"
        put(ws, row_number, 3, status, "text", "@")
        for index, period in enumerate(grid.forecasts, start=4):
            raw = driver.get(_source_period(period))
            visible = active[driver_id][period] if status == "Approved" else raw
            if visible in (None, ""):
                put(ws, row_number, index, None, "missing", FMT_M)
            else:
                try:
                    visible = float(visible)
                except (TypeError, ValueError) as exc:
                    raise EngineError(f"{driver_id} {_source_period(period)} is not numeric") from exc
                put(ws, row_number, index, visible,
                    "hardcode" if status == "Approved" else "proposed", FMT_M)
            references.bind(f"input.{driver_id}.{period}", ws.title, row_number, index)


def build_workbook(spec: Mapping[str, Any], actuals: Iterable[Mapping[str, Any]],
                   drivers: Iterable[Mapping[str, Any]], company_module: CompanyModule,
                   approval_manifest: Mapping[str, Any] | None = None) -> BuildContext:
    """Build the evidence/driver spine and render company-module outputs by metric key."""
    historical = tuple(spec["periods"]["historical_quarters"])
    grid = PeriodGrid(historical, _forecast_periods(spec))
    driver_rows = list(drivers)
    actual_map = _actual_map(actuals, historical)
    active = active_driver_values(driver_rows, grid.forecasts, approval_manifest,
                                  forecast_gate_approved=_forecast_gate(spec))
    workbook_spec = spec.get("workbook")
    sheets = tuple(workbook_spec.get("sheets", ())) if isinstance(workbook_spec, Mapping) else ()
    if not sheets or len(sheets) != len(set(sheets)):
        raise EngineError("model_spec.json workbook.sheets must be a non-empty unique list")
    nine_sheet_mode = set(sheets) == set(NINE_SHEET_ORDER)
    if nine_sheet_mode and sheets != NINE_SHEET_ORDER:
        raise EngineError("nine-sheet workbook sheets must use the agreed order")
    if nine_sheet_mode and len(grid.forecasts) != 8:
        raise EngineError("the nine-sheet workbook requires exactly eight forecast quarters")
    if not nine_sheet_mode and {"SourceData", "Drivers"} - set(sheets):
        raise EngineError("legacy model_spec.json workbook.sheets must include SourceData and Drivers")
    workbook = Workbook()
    workbook.active.title = sheets[0]
    for sheet in sheets[1:]:
        workbook.create_sheet(sheet)
    references = KeyedReferences(workbook)
    _write_actuals(workbook["SourceData"], actual_map, grid, references, spec.get("price"))
    module_rows = _module_rows(company_module, actual_map, active, grid)
    engine_sheets = {"SourceData", "Inputs"} if nine_sheet_mode else {"SourceData", "Drivers"}
    if not nine_sheet_mode:
        _write_drivers(workbook["Drivers"], driver_rows, active, grid, references)
    company_sheets = set(sheets) - engine_sheets
    if set(module_rows) != company_sheets:
        raise EngineError("company module sheets must exactly match model_spec.json workbook.sheets")
    forecast_active = any(value is not None for row in active.values() for value in row.values())
    if nine_sheet_mode:
        _write_inputs(workbook["Inputs"], driver_rows, active, grid, references)
    for sheet in sheets:
        if sheet in company_sheets:
            _write_module_rows(workbook[sheet], module_rows[sheet], grid, references,
                               forecast_active, legacy_order=not nine_sheet_mode)
    return BuildContext(workbook, grid, references, actual_map, active, module_rows)
