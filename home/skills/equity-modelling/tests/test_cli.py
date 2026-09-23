"""End-to-end contract tests for the artifact-driven Stage 3 CLI."""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from cli import build, check  # noqa: E402
from evidence import build_actuals, freeze, load_fact_map  # noqa: E402
from pull_fixture import create_source  # noqa: E402


MODULE = '''
class Module:
    def workbook_rows(self, actuals, drivers, periods):
        regional = {"revenue": {"2026Q1": "=em_actual_revenue_2026Q1", "2026Q2": "=em_actual_revenue_2026Q2", "2026Q3E": "=em_actual_revenue_2026Q2"}}
        return {
            "RegionalModel": regional,
            "Earnings": {"adjusted_eps": {"2026Q3E": "=em_model_RegionalModel_revenue_2026Q3E"}},
            "Scenarios": {"base": {"2026Q3E": "=em_model_Earnings_adjusted_eps_2026Q3E"}},
        }
company_module = Module()
'''


class CliTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.project = Path(self.temp.name)
        (self.project / "evidence").mkdir()
        spec, facts = create_source(self.project / "pull" / "data", "VRT")
        spec["periods"].update({"historical_quarters": ["2026Q1", "2026Q2"],
                                 "forecast_quarters": ["2026Q3"]})
        spec["price"] = {"value": 250.0, "date": "2026-09-18"}
        spec["workbook"] = {"filename": "VRT_2026-09-22_v2.xlsx",
                            "sheets": ["SourceData", "Drivers", "RegionalModel", "Earnings", "Scenarios"]}
        (self.project / "model_spec.json").write_text(json.dumps(spec))
        (self.project / "model_VRT.py").write_text(MODULE)
        fact_map = {"schema_version": 1, "ticker": "VRT", "facts": facts}
        (self.project / "fact_map.json").write_text(json.dumps(fact_map))
        mappings = load_fact_map(self.project / "fact_map.json")
        freeze(build_actuals(mappings, spec, ["2026Q1", "2026Q2"]), self.project / "evidence")
        benchmark = {
            "metric": "price_dated", "period": "2026-09-18", "value": "250.0",
            "units": "USD/share", "basis": "market", "source": "Yahoo",
            "as_of_date": "2026-09-18", "locator": next(
                f"raw_payload:yahoo/{item['sha256']}/payload.json#index={{date}}&field=Close"
                for item in spec["source_boundary"]["originals"] if item["provider"] == "yahoo"),
            "lineage_scheme": "raw_payload", "lineage_path": str(
                self.project / "pull" / "data" / "raw" / "VRT" / "yahoo" /
                next(item["sha256"] for item in spec["source_boundary"]["originals"]
                     if item["provider"] == "yahoo") / "payload.json"),
            "lineage_key": "index=2026-09-18T04:00:00.000Z;field=Close", "transform": "direct",
            "dimension": "", "notes": "",
        }
        with (self.project / "evidence" / "benchmarks.csv").open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=benchmark.keys())
            writer.writeheader(); writer.writerow(benchmark)
        with (self.project / "drivers.csv").open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["driver_id", "driver_name", "status", "2026Q3"])
            writer.writeheader()
            writer.writerow({"driver_id": "revenue_growth", "driver_name": "Revenue growth",
                             "status": "proposed", "2026Q3": "0.2"})

    def tearDown(self):
        self.temp.cleanup()

    def test_builds_once_from_frozen_artifacts_and_keeps_forecast_blank(self):
        output = build(self.project)
        self.assertTrue(output.is_file())
        report = check(self.project)
        self.assertTrue(report.ok, report.failures)
        from openpyxl import load_workbook
        workbook = load_workbook(output, data_only=False)
        self.assertIsNone(workbook["RegionalModel"]["E7"].value)
        self.assertIsNone(workbook["SourceData"]["B4"].value)
        with self.assertRaises(FileExistsError):
            build(self.project)

    def test_full_reports_replay_coverage_for_actuals_and_benchmarks(self):
        build(self.project)
        report = check(self.project, mode="full")
        self.assertTrue(report.ok, report.failures)
        self.assertIn("source replay: 3 actuals, 1 benchmarks", report.coverage)

    def test_build_rejects_tampered_frozen_value(self):
        rows, fieldnames = self._actuals()
        rows[0]["value"] = "999"
        self._write_actuals(rows, fieldnames)
        self._assert_build_rejects("revenue 2026Q1: frozen value")

    def test_full_check_rejects_tampered_actual(self):
        build(self.project)
        rows, fields = self._actuals()
        rows[0]["value"] = "999"
        self._write_actuals(rows, fields)
        report = check(self.project, mode="full")
        self.assertFalse(report.ok)
        self.assertTrue(any("revenue 2026Q1" in failure and "does not replay" in failure
                            for failure in report.failures))

    def test_full_check_rejects_tampered_benchmark(self):
        build(self.project)
        path = self.project / "evidence" / "benchmarks.csv"
        with path.open(newline="") as handle:
            reader = csv.DictReader(handle); rows = list(reader); fields = reader.fieldnames
        rows[0]["value"] = "999"
        with path.open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader(); writer.writerows(rows)
        report = check(self.project, mode="full")
        self.assertFalse(report.ok)
        self.assertTrue(any("benchmark price_dated 2026-09-18" in failure for failure in report.failures))

    def test_build_rejects_price_missing_from_benchmarks(self):
        path = self.project / "evidence" / "benchmarks.csv"
        with path.open("w", newline="") as handle:
            handle.write("metric,period,value\\n")
        self._assert_build_rejects("expected exactly one replayed price_dated benchmark, found 0")

    def test_build_rejects_price_value_conflicting_with_replayed_benchmark(self):
        spec_path = self.project / "model_spec.json"
        spec = json.loads(spec_path.read_text())
        spec["price"]["value"] = 251.0
        spec_path.write_text(json.dumps(spec))
        self._assert_build_rejects("value does not match replayed price_dated benchmark")

    def test_build_rejects_duplicate_dated_price_benchmarks(self):
        path = self.project / "evidence" / "benchmarks.csv"
        with path.open(newline="") as handle:
            reader = csv.DictReader(handle); rows = list(reader); fields = reader.fieldnames
        rows.append(dict(rows[0]))
        with path.open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader(); writer.writerows(rows)
        self._assert_build_rejects("expected exactly one replayed price_dated benchmark, found 2")

    def test_build_rejects_tampered_benchmark(self):
        path = self.project / "evidence" / "benchmarks.csv"
        with path.open(newline="") as handle:
            reader = csv.DictReader(handle); rows = list(reader); fields = reader.fieldnames
        rows[0]["value"] = "999"
        with path.open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader(); writer.writerows(rows)
        self._assert_build_rejects("benchmark price_dated 2026-09-18: frozen value does not replay")

    def test_build_rejects_changed_pinned_table(self):
        spec = json.loads((self.project / "model_spec.json").read_text())
        table = spec["source_boundary"]["tables"][0]
        path = (self.project / "pull" / "data" / "tables" / table["ticker"] /
                table["run_id"] / f"{table['name']}.parquet")
        path.write_bytes(path.read_bytes() + b"tamper")
        with self.assertRaisesRegex(ValueError, "pinned table hash mismatch") as raised:
            build(self.project)
        self.assertIn("revenue 2026Q1", str(raised.exception))

    def test_full_check_rejects_changed_pinned_raw_source(self):
        output = build(self.project)
        source = next(item for item in json.loads((self.project / "model_spec.json").read_text())
                      ["source_boundary"]["originals"] if item["provider"] == "sec")
        raw = self.project / "pull" / "data" / "raw" / "VRT" / "sec" / source["sha256"] / source["filename"]
        raw.write_bytes(raw.read_bytes() + b"<!-- changed -->")
        report = check(self.project, mode="full")
        self.assertFalse(report.ok)
        self.assertTrue(any("adjusted_operating_profit 2026Q2" in failure
                            and "pinned original hash mismatch" in failure for failure in report.failures))

    def test_build_rejects_missing_source_boundary_metadata(self):
        path = self.project / "model_spec.json"
        spec = json.loads(path.read_text())
        del spec["source_boundary"]["tables"]
        path.write_text(json.dumps(spec))
        self._assert_build_rejects("source_boundary")

    def test_full_rejects_positional_cross_sheet_formula(self):
        output = build(self.project)
        from openpyxl import load_workbook
        workbook = load_workbook(output)
        workbook["RegionalModel"]["E7"] = "='SourceData'!C7"
        workbook.save(output)
        report = check(self.project, mode="full")
        self.assertFalse(report.ok)
        self.assertIn("RegionalModel!E7 uses a positional cross-sheet reference", report.failures)

    def _assert_build_rejects(self, expected: str) -> None:
        with self.assertRaises(ValueError) as raised:
            build(self.project)
        self.assertIn(expected, str(raised.exception))

    def _actuals(self) -> tuple[list[dict[str, str]], list[str]]:
        path = self.project / "evidence" / "actuals.csv"
        with path.open(newline="") as handle:
            reader = csv.DictReader(handle)
            return list(reader), list(reader.fieldnames or [])

    def _write_actuals(self, rows: list[dict[str, str]], fieldnames: list[str]) -> None:
        with (self.project / "evidence" / "actuals.csv").open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

    def test_build_rejects_missing_fact_map(self):
        (self.project / "fact_map.json").unlink()
        self._assert_build_rejects("missing fact_map.json")

    def test_build_rejects_non_object_fact_map(self):
        (self.project / "fact_map.json").write_text(json.dumps(["not", "an", "object"]))
        self._assert_build_rejects("fact_map.json")

    def test_build_rejects_actuals_missing_lineage_columns(self):
        rows, _ = self._actuals()
        reduced = [{k: row[k] for k in ("metric", "period", "value")} for row in rows]
        self._write_actuals(reduced, ["metric", "period", "value"])
        self._assert_build_rejects("actuals.csv is missing columns")

    def test_build_rejects_populated_value_without_lineage(self):
        rows, fieldnames = self._actuals()
        rows[0]["lineage_path"] = ""
        self._write_actuals(rows, fieldnames)
        self._assert_build_rejects("populated value lacks lineage")

    def test_build_rejects_unavailable_value_without_reason(self):
        rows, fieldnames = self._actuals()
        rows[0]["value"] = ""
        rows[0]["provenance_status"] = "verified"
        rows[0]["notes"] = ""
        self._write_actuals(rows, fieldnames)
        self._assert_build_rejects("unavailable value has no specific missing reason")

    def test_build_rejects_bare_status_without_specific_notes(self):
        rows, fieldnames = self._actuals()
        rows[0]["value"] = ""
        rows[0]["provenance_status"] = "unavailable"
        rows[0]["notes"] = "unavailable"
        self._write_actuals(rows, fieldnames)
        self._assert_build_rejects("unavailable value has no specific missing reason")

    def test_build_accepts_unavailable_value_with_specific_reason(self):
        rows, fieldnames = self._actuals()
        rows[0]["value"] = ""
        rows[0]["provenance_status"] = "unavailable"
        rows[0]["lineage_path"] = ""
        rows[0]["notes"] = "issuer did not separately disclose the balance"
        self._write_actuals(rows, fieldnames)
        self.assertTrue(build(self.project).is_file())

    def test_build_rejects_missing_expected_actual_row(self):
        rows, fieldnames = self._actuals()
        self._write_actuals([row for row in rows if row["period"] != "2026Q2"], fieldnames)
        self._assert_build_rejects(
            "fact_map.json declares revenue 2026Q2 but actuals.csv has no row")

    def test_build_rejects_duplicate_actuals_rows(self):
        rows, fieldnames = self._actuals()
        self._write_actuals(rows + [dict(rows[0])], fieldnames)
        self._assert_build_rejects("duplicate row")


if __name__ == "__main__":
    unittest.main()
