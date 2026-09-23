"""Contract tests for generic, project-level driver rollover."""
from __future__ import annotations

import csv
import hashlib
import json
import shutil
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))
from rollover import RolloverError, rollover  # noqa: E402

MODULE = '''
class Module:
    def workbook_rows(self, actuals, drivers, periods):
        return {}
company_module = Module()
'''

sys.path.insert(0, str(Path(__file__).parent))
from evidence import build_actuals, freeze  # noqa: E402
from pull_fixture import create_source, fact_mappings, price_benchmark_row  # noqa: E402

DRIVER_COLUMNS = ("driver_id", "driver_name", "unit", "status", "approved_by", "approved_on",
                  "approval_id")


def _write_project(root: Path, forecast: list[str], drivers: list[dict[str, str]],
                   *, manifest: dict | None = None, filename: str = "TST_model.xlsx") -> Path:
    (root / "evidence").mkdir(parents=True, exist_ok=True)
    source_root = root / "pull" / "data"
    if source_root.exists():
        shutil.rmtree(source_root)
    spec, facts = create_source(source_root, "TST")
    spec.update({
        "periods": {"historical_quarters": ["2026Q1", "2026Q2"], "forecast_quarters": forecast},
        "forecast_gate": {"approved": False},
        "price": {"value": 250.0, "date": "2026-09-18"},
        "workbook": {"filename": filename, "sheets": ["SourceData", "Drivers"]},
    })
    (root / "model_spec.json").write_text(json.dumps(spec))
    (root / "model_TST.py").write_text(MODULE)
    (root / "fact_map.json").write_text(json.dumps({"schema_version": 1, "ticker": "TST", "facts": facts}))
    frozen = build_actuals(fact_mappings(facts), spec, ["2026Q1", "2026Q2"])
    freeze(frozen, root / "evidence")
    benchmark = price_benchmark_row(spec)
    with (root / "evidence" / "benchmarks.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(benchmark))
        writer.writeheader()
        writer.writerow(benchmark)
    fieldnames = list(DRIVER_COLUMNS) + list(forecast)
    with (root / "drivers.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for driver in drivers:
            writer.writerow(driver)
    if manifest is not None:
        (root / "drivers.approval.json").write_text(json.dumps(manifest))
    return root


def _read_drivers(project: Path) -> list[dict[str, str]]:
    with (project / "drivers.csv").open(newline="") as handle:
        return list(csv.DictReader(handle))


def _by_id(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    return {str(row["driver_id"]).strip(): row for row in rows}


def _tree_hashes(root: Path) -> dict[str, str]:
    hashes: dict[str, str] = {}
    for path in sorted(root.rglob("*")):
        if path.is_file() and "__pycache__" not in path.parts:
            hashes[path.relative_to(root).as_posix()] = hashlib.sha256(path.read_bytes()).hexdigest()
    return hashes


class RolloverTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        base = Path(self.temp.name)
        self.prior = base / "prior"
        self.new = base / "new"
        self.output = base / "rolled"

        _write_project(self.prior, ["2026Q3", "2026Q4", "2027Q1"], [
            {"driver_id": "growth", "driver_name": "Revenue growth", "unit": "percent",
             "status": "approved", "approved_by": "elias", "approved_on": "2026-09-22",
             "approval_id": "APR-1", "2026Q3": "0.4", "2026Q4": "0.5", "2027Q1": "0.33"},
            {"driver_id": "margin", "driver_name": "Adjusted operating margin", "unit": "percent",
             "status": "proposed", "approved_by": "", "approved_on": "", "approval_id": "",
             "2026Q3": "0.25", "2026Q4": "0.26", "2027Q1": "0.23"},
            {"driver_id": "legacy", "driver_name": "Retired driver", "unit": "USD millions",
             "status": "proposed", "approved_by": "", "approved_on": "", "approval_id": "",
             "2026Q3": "1", "2026Q4": "2", "2027Q1": "3"},
        ], manifest={"approvals": {"APR-1": {"sha256": "a" * 64}}})

        _write_project(self.new, ["2026Q4", "2027Q1", "2027Q2"], [
            {"driver_id": "growth", "driver_name": "Revenue growth", "unit": "percent",
             "status": "approved", "approved_by": "elias", "approved_on": "2026-09-22",
             "approval_id": "APR-2", "2026Q4": "", "2027Q1": "", "2027Q2": ""},
            {"driver_id": "margin", "driver_name": "Renamed margin", "unit": "percent",
             "status": "proposed", "approved_by": "", "approved_on": "", "approval_id": "",
             "2026Q4": "", "2027Q1": "", "2027Q2": ""},
            {"driver_id": "extension", "driver_name": "New extension driver", "unit": "percent",
             "status": "proposed", "approved_by": "", "approved_on": "", "approval_id": "",
             "2026Q4": "", "2027Q1": "", "2027Q2": ""},
        ], manifest={"approvals": {"APR-2": {"sha256": "b" * 64}}})

    def tearDown(self):
        self.temp.cleanup()

    def test_carries_only_overlapping_period_values_by_driver_id(self):
        report = rollover(self.prior, self.new, self.output)
        growth = _by_id(_read_drivers(self.output))["growth"]
        self.assertEqual(growth["2026Q4"], "0.5")
        self.assertEqual(growth["2027Q1"], "0.33")
        self.assertEqual(growth["2027Q2"], "")  # new-only period is never carried
        carried = {entry["driver_id"]: entry["periods"] for entry in report["carried"]}
        self.assertEqual(carried["growth"], {"2026Q4": "0.5", "2027Q1": "0.33"})
        self.assertEqual(report["overlapping_periods"], ["2026Q4", "2027Q1"])

    def test_preserves_a_deliberate_blank_as_blank(self):
        _write_project(self.new, ["2026Q4", "2027Q1", "2027Q2"], [
            {"driver_id": "growth", "driver_name": "Revenue growth", "unit": "percent",
             "status": "proposed", "approved_by": "", "approved_on": "", "approval_id": "",
             "2026Q4": "", "2027Q1": "", "2027Q2": ""},
            {"driver_id": "margin", "driver_name": "Renamed margin", "unit": "percent",
             "status": "proposed", "approved_by": "", "approved_on": "", "approval_id": "",
             "2026Q4": "", "2027Q1": "", "2027Q2": ""},
            {"driver_id": "extension", "driver_name": "New extension driver", "unit": "percent",
             "status": "proposed", "approved_by": "", "approved_on": "", "approval_id": "",
             "2026Q4": "", "2027Q1": "", "2027Q2": ""},
        ])
        # Prior `growth` 2027Q1 is 0.33; blank it to prove a blank prior value survives as blank.
        prior_rows = _read_drivers(self.prior)
        for row in prior_rows:
            if row["driver_id"] == "growth":
                row["2027Q1"] = ""
        with (self.prior / "drivers.csv").open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(prior_rows[0].keys()))
            writer.writeheader()
            writer.writerows(prior_rows)
        rollover(self.prior, self.new, self.output)
        growth = _by_id(_read_drivers(self.output))["growth"]
        self.assertEqual(growth["2027Q1"], "")

    def test_reports_retired_period_and_drops_its_column(self):
        report = rollover(self.prior, self.new, self.output)
        self.assertEqual(report["retired_periods"], ["2026Q3"])
        with (self.output / "drivers.csv").open(newline="") as handle:
            self.assertNotIn("2026Q3", csv.DictReader(handle).fieldnames)

    def test_withholds_changed_definition_and_reports_removed_driver(self):
        report = rollover(self.prior, self.new, self.output)
        margin = _by_id(_read_drivers(self.output))["margin"]
        self.assertEqual(margin["2026Q4"], "")  # prior 0.26 must not be copied across a rename
        self.assertEqual(margin["2027Q1"], "")
        self.assertEqual([entry["driver_id"] for entry in report["withheld"]], ["margin"])
        self.assertEqual(report["removed_drivers"], ["legacy"])
        self.assertNotIn("legacy", _by_id(_read_drivers(self.output)))
        self.assertEqual(report["added_drivers"], ["extension"])

    def test_every_row_is_proposed_and_no_approval_is_carried(self):
        report = rollover(self.prior, self.new, self.output)
        self.assertEqual(report["approval"],
                         {"carried": False, "manifest_copied": False, "statuses": "proposed"})
        for row in _read_drivers(self.output):
            self.assertEqual(row["status"], "proposed")
            self.assertEqual(row["approved_by"], "")
            self.assertEqual(row["approved_on"], "")
            self.assertEqual(row["approval_id"], "")
        self.assertFalse((self.output / "drivers.approval.json").exists())

    def test_prior_and_new_projects_remain_byte_identical(self):
        prior_hashes = _tree_hashes(self.prior)
        new_hashes = _tree_hashes(self.new)
        rollover(self.prior, self.new, self.output)
        self.assertEqual(_tree_hashes(self.prior), prior_hashes)
        self.assertEqual(_tree_hashes(self.new), new_hashes)

    def test_build_succeeds_but_forecast_output_stays_blank(self):
        report = rollover(self.prior, self.new, self.output)
        delivery = report["delivery"]
        self.assertTrue(delivery["ok"], delivery["failures"])
        self.assertTrue(Path(delivery["workbook"]).is_file())
        from openpyxl import load_workbook
        workbook = load_workbook(delivery["workbook"], data_only=True)
        source_values = {cell.value for row in workbook["SourceData"].iter_rows() for cell in row}
        self.assertIn(100, source_values)  # actuals still load from the pinned synthetic pull
        drivers_sheet = workbook["Drivers"]
        for row in range(7, drivers_sheet.max_row + 1):
            for column in range(4, drivers_sheet.max_column + 1):
                self.assertIsNone(drivers_sheet.cell(row, column).value)

    def test_missing_status_column_is_added_as_proposed(self):
        rows = [
            {"driver_id": "growth", "driver_name": "Revenue growth", "unit": "percent",
             "2026Q4": "", "2027Q1": "", "2027Q2": ""},
            {"driver_id": "extension", "driver_name": "New extension driver", "unit": "percent",
             "2026Q4": "", "2027Q1": "", "2027Q2": ""},
        ]
        with (self.new / "drivers.csv").open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        report = rollover(self.prior, self.new, self.output)
        output_rows = _read_drivers(self.output)
        self.assertIn("status", output_rows[0])
        self.assertTrue(all(row["status"] == "proposed" for row in output_rows))
        self.assertTrue(report["delivery"]["ok"], report["delivery"]["failures"])

    def test_duplicate_driver_ids_fail(self):
        rows = _read_drivers(self.new)
        rows.append(dict(rows[0]))
        with (self.new / "drivers.csv").open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        with self.assertRaisesRegex(RolloverError, "duplicate driver_id"):
            rollover(self.prior, self.new, self.output)

    def test_refuses_to_overwrite_existing_output(self):
        self.output.mkdir()
        (self.output / "keep.txt").write_text("untouched")
        with self.assertRaisesRegex(RolloverError, "refusing to overwrite existing output"):
            rollover(self.prior, self.new, self.output)
        self.assertTrue((self.output / "keep.txt").is_file())

    def test_cli_rollover_command_writes_output_and_report(self):
        report_path = Path(self.temp.name) / "report.json"
        result = subprocess.run(
            [sys.executable, str(SCRIPTS / "cli.py"), str(self.new), "rollover",
             "--prior-project", str(self.prior), "--output-project", str(self.output),
             "--report", str(report_path)],
            capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((self.output / "drivers.csv").is_file())
        payload = json.loads(report_path.read_text())
        self.assertEqual(payload["retired_periods"], ["2026Q3"])
        self.assertTrue(payload["delivery"]["ok"])


if __name__ == "__main__":
    unittest.main()
