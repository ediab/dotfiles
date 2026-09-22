"""Checks for the override extract/restore helper.

Run: python3 -m unittest discover -s tests -v   (from the skill directory)
"""
from __future__ import annotations

import hashlib
import sys
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import overrides as O  # noqa: E402

PERIODS = ["Q3 2026E", "Q4 2026E", "Q1 2027E"]
DRIVERS = ["Revenue growth", "Adjusted margin"]


def make_workbook(path: Path, periods=PERIODS, drivers=DRIVERS, first_col=3,
                  first_block_row=8, overrides=None, definition=None, header_row=6):
    """Build a contract-shaped override sheet.  ``overrides`` maps (driver, period) -> value."""
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Assumptions"
    for offset, label in enumerate(periods):
        ws.cell(header_row, first_col + offset, label)
    for index, driver in enumerate(drivers):
        row = first_block_row + index * 4
        ws.cell(row, 2, f"{driver} | consensus default")
        ws.cell(row + 1, 2, f"{driver} | YOUR override")
        ws.cell(row + 2, 2, f"{driver} | active input")
        for offset in range(len(periods)):
            col = first_col + offset
            letter = ws.cell(header_row, col).column_letter
            ws.cell(row, col, 0.10 + index * 0.01)
            ws.cell(row + 2, col, definition(row, col, letter) if definition else
                    f'=IF({letter}{row + 1}<>"",{letter}{row + 1},{letter}{row})')
    for (driver, period), value in (overrides or {}).items():
        row = first_block_row + drivers.index(driver) * 4 + 1
        ws.cell(row, first_col + periods.index(period), value)
    workbook.save(path)
    return path


class OverrideTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def digest(self, path):
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()

    # ---------------------------------------------------------------- extraction
    def test_extract_records_identity_not_coordinates(self):
        src = make_workbook(self.tmp / "src.xlsx",
                            overrides={("Revenue growth", "Q1 2027E"): 0.25,
                                       ("Adjusted margin", "Q3 2026E"): 0.0})
        found = O.extract(src)
        self.assertEqual(len(found.overrides), 2)
        self.assertEqual(found.get("Revenue growth", "Q1 2027E").value, 0.25)
        self.assertEqual(found.get("Revenue growth", "2027Q1E").value, 0.25)   # canonical accepted
        self.assertEqual(found.drivers["Adjusted margin"].override_row, 13)
        self.assertEqual(sorted(found.periods.values()), ["2026Q3E", "2026Q4E", "2027Q1E"])

    def test_zero_is_a_valid_override_and_blank_is_not(self):
        src = make_workbook(self.tmp / "src.xlsx",
                            overrides={("Adjusted margin", "Q4 2026E"): 0})
        found = O.extract(src)
        self.assertEqual(len(found.overrides), 1)
        self.assertEqual(found.get("Adjusted margin", "Q4 2026E").value, 0.0)
        self.assertIsNone(found.get("Revenue growth", "Q4 2026E"))

    # ------------------------------------------------------------- restore paths
    def test_restore_follows_moved_rows_and_columns(self):
        original = make_workbook(self.tmp / "original.xlsx",
                                 overrides={("Revenue growth", "Q1 2027E"): 0.25})
        found = O.extract(original)
        before = self.digest(original)

        # The rebuild moved the blocks six rows down and the periods two columns right.
        template = make_workbook(self.tmp / "rebuild.xlsx", first_col=5, first_block_row=14)
        out = self.tmp / "updated.xlsx"
        report = O.restore(template, out, found)

        self.assertEqual(len(report.restored), 1)
        self.assertEqual(report.definition_changed, [])
        self.assertEqual(self.digest(original), before)
        self.assertEqual(self.digest(template), self.digest(self.tmp / "rebuild.xlsx"))
        ws = load_workbook(out)["Assumptions"]
        self.assertEqual(ws.cell(15, 7).value, 0.25)          # Q1 2027E is column G here
        self.assertIsNone(ws.cell(15, 5).value)

    def test_rollover_retires_the_quarter_that_became_actual(self):
        original = make_workbook(self.tmp / "original.xlsx",
                                 overrides={("Revenue growth", "Q3 2026E"): 0.30,
                                            ("Revenue growth", "Q4 2026E"): 0.35})
        found = O.extract(original)
        before = self.digest(original)

        # One quarter later: Q3 2026E is now an actual quarter and left the grid.
        template = make_workbook(self.tmp / "rebuild.xlsx",
                                 periods=["Q4 2026E", "Q1 2027E", "Q2 2027E"])
        out = self.tmp / "updated.xlsx"
        report = O.restore(template, out, found)

        self.assertEqual([item.period for item in report.retired_period], ["2026Q3E"])
        self.assertEqual([item.period for item in report.restored], ["2026Q4E"])
        self.assertEqual(self.digest(original), before)
        ws = load_workbook(out)["Assumptions"]
        self.assertEqual(ws.cell(9, 3).value, 0.35)           # still attached to Q4 2026E
        self.assertIsNone(ws.cell(9, 4).value)

    def test_cleared_override_leaves_the_default(self):
        original = make_workbook(self.tmp / "original.xlsx",
                                 overrides={("Revenue growth", "Q4 2026E"): 0.35})
        cleared = make_workbook(self.tmp / "cleared.xlsx")     # user erased every override cell
        self.assertEqual(O.extract(cleared).overrides, {})
        self.assertEqual(O.extract(original).get("Revenue growth", "Q4 2026E").value, 0.35)

    def test_removed_driver_is_reported_and_never_migrated(self):
        original = make_workbook(self.tmp / "original.xlsx",
                                 overrides={("Adjusted margin", "Q3 2026E"): 0.27})
        found = O.extract(original)

        dropped = make_workbook(self.tmp / "dropped.xlsx", drivers=["Revenue growth", "New driver"])
        report = O.restore(dropped, self.tmp / "out1.xlsx", found)
        self.assertEqual([item.driver for item in report.removed_driver], ["Adjusted margin"])
        self.assertEqual(report.restored, [])
        self.assertIsNone(load_workbook(self.tmp / "out1.xlsx")["Assumptions"].cell(9, 3).value)

        renamed = make_workbook(self.tmp / "renamed.xlsx",
                                drivers=["Revenue growth", "Adjusted operating margin"])
        report = O.restore(renamed, self.tmp / "out2.xlsx", found)
        self.assertEqual([item.driver for item in report.removed_driver], ["Adjusted margin"])
        self.assertIsNone(load_workbook(self.tmp / "out2.xlsx")["Assumptions"].cell(13, 3).value)

    def test_changed_definition_is_withheld_until_allowed(self):
        original = make_workbook(self.tmp / "original.xlsx",
                                 overrides={("Revenue growth", "Q3 2026E"): 0.30})
        found = O.extract(original)

        def different(row, col, letter):
            return f'=MAX(0,IF({letter}{row + 1}<>"",{letter}{row + 1},{letter}{row}))'

        template = make_workbook(self.tmp / "rebuild.xlsx", definition=different)
        out = self.tmp / "updated.xlsx"
        report = O.restore(template, out, found)
        self.assertEqual([item.driver for item in report.definition_changed], ["Revenue growth"])
        self.assertEqual(report.restored, [])
        self.assertIsNone(load_workbook(out)["Assumptions"].cell(9, 3).value)

        forced = make_workbook(self.tmp / "forced.xlsx", definition=different)
        report = O.restore(forced, self.tmp / "forced-out.xlsx", found, allow_definition_change=True)
        self.assertEqual(len(report.restored), 1)
        self.assertEqual(report.definition_changed, [])
        self.assertEqual(load_workbook(self.tmp / "forced-out.xlsx")["Assumptions"].cell(9, 3).value, 0.30)

    # --------------------------------------------------------- safety and errors
    def test_originals_are_never_written(self):
        original = make_workbook(self.tmp / "original.xlsx",
                                 overrides={("Revenue growth", "Q3 2026E"): 0.30})
        found = O.extract(original)
        before = self.digest(original)

        template = make_workbook(self.tmp / "rebuild.xlsx")
        template_before = self.digest(template)
        report = O.restore(template, self.tmp / "updated.xlsx", found)

        self.assertTrue(report.original_unchanged)
        self.assertEqual(self.digest(original), before)
        self.assertEqual(self.digest(template), template_before)
        with self.assertRaises(O.OverrideError):
            O.restore(template, template, found)
        self.assertEqual(self.digest(template), template_before)

    def test_invalid_override_values_are_rejected(self):
        for bad in ("=1+1", "text", True):
            path = make_workbook(self.tmp / "bad.xlsx")
            workbook = load_workbook(path)
            workbook["Assumptions"].cell(9, 3, bad)
            workbook.save(path)
            with self.assertRaises(O.OverrideError, msg=f"{bad!r} should be rejected"):
                O.extract(path)

    def test_incompatible_metadata_is_rejected(self):
        def broken(mutate):
            path = make_workbook(self.tmp / "broken.xlsx")
            workbook = load_workbook(path)
            mutate(workbook["Assumptions"])
            workbook.save(path)
            return path

        with self.assertRaises(O.OverrideError):     # unparseable period label
            O.extract(broken(lambda ws: ws.cell(6, 3, "next quarter")))
        with self.assertRaises(O.OverrideError):     # duplicate period labels
            O.extract(broken(lambda ws: ws.cell(6, 4, "Q3 2026E")))
        with self.assertRaises(O.OverrideError):     # block without a default row above
            O.extract(broken(lambda ws: ws.cell(8, 2, "Revenue growth (no suffix)")))
        with self.assertRaises(O.OverrideError):     # duplicate override row for one driver
            O.extract(broken(lambda ws: ws.cell(20, 2, "Revenue growth | YOUR override")))
        with self.assertRaises(O.OverrideError):     # no period labels at all
            O.extract(broken(lambda ws: [setattr(ws.cell(6, col), "value", None) for col in (3, 4, 5)]))
        with self.assertRaises(O.OverrideError):     # override sheet missing
            O.extract(broken(lambda ws: setattr(ws.parent["Assumptions"], "title", "Drivers")))

    def test_restore_rejects_missing_sheet_and_duplicate_driver_rows(self):
        original = make_workbook(self.tmp / "original.xlsx",
                                 overrides={("Revenue growth", "Q3 2026E"): 0.30})
        found = O.extract(original)

        template = make_workbook(self.tmp / "rebuild.xlsx")
        workbook = load_workbook(template)
        workbook["Assumptions"].title = "Drivers"
        workbook.save(template)
        with self.assertRaises(O.OverrideError):
            O.restore(template, self.tmp / "out.xlsx", found)

        duplicate = make_workbook(self.tmp / "duplicate.xlsx")
        workbook = load_workbook(duplicate)
        workbook["Assumptions"].cell(20, 2, "Revenue growth | YOUR override")
        workbook["Assumptions"].cell(21, 2, "Revenue growth | active input")
        workbook.save(duplicate)
        with self.assertRaises(O.OverrideError):
            O.restore(duplicate, self.tmp / "out2.xlsx", found)

    def test_json_round_trip(self):
        original = make_workbook(self.tmp / "original.xlsx",
                                 overrides={("Adjusted margin", "Q1 2027E"): 0.231})
        O.extract(original).save(self.tmp / "overrides.json")
        loaded = O.OverrideSet.load(self.tmp / "overrides.json")
        self.assertEqual(loaded.get("Adjusted margin", "Q1 2027E").value, 0.231)

        template = make_workbook(self.tmp / "rebuild.xlsx", periods=["Q4 2026E", "Q1 2027E"])
        report = O.restore(template, self.tmp / "updated.xlsx", loaded)
        self.assertEqual(len(report.restored), 1)
        self.assertEqual(load_workbook(self.tmp / "updated.xlsx")["Assumptions"].cell(13, 4).value, 0.231)


if __name__ == "__main__":
    unittest.main()
